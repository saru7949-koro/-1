"""パイプライン不足分析(代理店軸)。作業/snapshots/<日付>/ のコピーだけを読む。

出力: 出力/パイプライン不足分析_<日付>.xlsx / .md
顧客名(施設)・個人名・自由記述の原文は出力しない。案件はスナップショットの Excel 行番号で参照する。
"""
import math
import re
import sys
import tomllib
from pathlib import Path

import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
from load import load_pipeline  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent

# LST MOC FY26 Sep 2026 slide 2 の記載値(照合用)
PLAN_MODELS = ("SH", "MA", "SA", "ID", "FP")  # 9F目標のあるモデル(CGXは対象外)
DECK_9F = pd.DataFrame(
    {"FCST案件": [19, 15, 3, 5, 7], "Backup": [32, 16, 13, 12, 15], "着地理論値": [21, 11, 6, 4, 7],
     "9F不足": [-4, -4, 0, -2, -1], "必要リード": [6, 6, 0, 7, 2]},
    index=["SH", "MA", "SA", "ID", "FP"])

HIGH = {"ID", "FP"}


# ---------- 共通 ----------

def norm_agent_token(t, alias):
    t = re.sub(r"[\s　]+", "", str(t))
    if not t or t.lower() == "nan":
        return None
    if t in alias:
        return alias[t]
    for k, v in alias.items():
        if len(k) >= 3 and t.startswith(k):
            return v
    return t


def split_agent(raw, cfg):
    """代理店文字列 → (一次店, 二次店)。表記は「二次店/一次店」が多いが逆順もあるため、
    主要一次店リストに該当するものを一次店とし、無ければ末尾を一次店とみなす。"""
    alias, primaries = cfg["agent_alias"], cfg["rules"]["primary_agents"]
    if raw is None or (isinstance(raw, float) and math.isnan(raw)) or str(raw).strip() == "":
        return "(未記入)", ""
    toks = []
    for t in re.split(r"[/／]", str(raw)):
        n = norm_agent_token(t, alias)
        if n is None:
            continue
        for p in primaries:  # 「池田理化から連絡 調整中」等の付記を除去
            if n.startswith(p) and n != p:
                n = p
        toks.append(n)
    if not toks:
        return "(未記入)", ""
    prim = next((p for p in primaries if p in toks), None)
    if prim is None:
        real = [t for t in toks if t != "未定"]
        prim = real[-1] if (real and len(toks) == len(real)) else "未定"
        if len(toks) == 1 and toks[0] != "未定":
            prim = toks[0]
    rest = [t for t in toks if t not in (prim, "未定")]
    return prim, "/".join(rest)


def issue_tags(text, kw):
    if not isinstance(text, str) or not text.strip():
        return []
    return [k for k, words in kw.items() if any(w.lower() in text.lower() for w in words)]


def to_date(v):
    if isinstance(v, (pd.Timestamp,)) or hasattr(v, "year"):
        return pd.Timestamp(v)
    return pd.NaT


def md_table(df, index=False):
    d = df.reset_index() if index else df
    cols = [str(c) for c in d.columns]
    out = ["| " + " | ".join(cols) + " |", "|" + "|".join("---" for _ in cols) + "|"]
    for r in d.itertuples(index=False):
        out.append("| " + " | ".join(_fmt(v) for v in r) + " |")
    return "\n".join(out)


def _fmt(v):
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return ""
    if isinstance(v, float):
        return f"{v:.1f}" if abs(v) >= 1 or v == 0 else f"{v:.2f}"
    return str(v).replace("|", "/").replace("\n", " ")


# ---------- 状態付け ----------

def classify(df, cfg):
    p = cfg["period"]
    d = df.copy()
    d["成約"] = d["区分"].str.match(r"^\d+FCT$")
    d["今期"] = d["時期"].between(p["fy_start"], p["fy_end"])
    d["open_FCT"] = d["今期"] & (d["区分"] == "FCT") & (d["案件確度"] > 0)
    d["open_BK"] = d["今期"] & (d["区分"] == "backup") & (d["案件確度"] > 0)
    wr = {m: v["win_rate"] for m, v in cfg["models"].items()}
    d["勝率"] = d["モデルイニシャル"].map(wr)
    bk = cfg["backup_rate"]
    in_plan = d["勝率"].notna()
    d["着地寄与"] = 0.0
    d.loc[d["今期"] & d["成約"] & in_plan, "着地寄与"] = d["台数"]
    d.loc[d["open_FCT"] & in_plan, "着地寄与"] = d["台数"] * d["勝率"]
    d.loc[d["open_BK"] & in_plan, "着地寄与"] = d["台数"] * bk
    return d


def landing(d, cfg, label):
    rows = []
    for m, mv in cfg["models"].items():
        x = d[d["モデルイニシャル"] == m]
        act = x.loc[x["今期"] & x["成約"], "台数"].sum()
        fct = x.loc[x["open_FCT"], "台数"].sum()
        bk = x.loc[x["open_BK"], "台数"].sum()
        land = act + fct * mv["win_rate"] + bk * cfg["backup_rate"]
        gap = land - mv["target_9f"]
        need = max(0.0, -gap) / mv["win_rate"]
        rows.append({"モデル": m, "BGT": mv["bgt"], "9F目標": mv["target_9f"], "実績(成約)": act,
                     "FCT案件": fct, "勝率": mv["win_rate"], "Backup": bk, "着地理論値": round(land, 2),
                     "不足(台)": round(gap, 2), "必要リード(件)": round(need, 2)})
    t = pd.DataFrame(rows)
    tot = t[["BGT", "9F目標", "実績(成約)", "FCT案件", "Backup", "着地理論値", "不足(台)", "必要リード(件)"]].sum()
    t.loc[len(t)] = {"モデル": "計", **tot.round(2).to_dict(), "勝率": None}
    t.insert(0, "データ", label)
    return t


# ---------- 9F → 10/5 の動き ----------

def movement(a, b, cfg):
    p = cfg["period"]
    for d in (a, b):
        d["key"] = (d["施設"].astype(str).str.replace(r"\s", "", regex=True) + "|"
                    + d["モデル詳細"].astype(str).str.strip())
        d["key"] = d["key"] + "#" + d.groupby("key").cumcount().astype(str)
    m = a.merge(b, on="key", how="outer", suffixes=("_9F", ""), indicator=True)
    infy = lambda s: s.between(p["fy_start"], p["fy_end"])  # noqa: E731
    base = m[infy(m["時期_9F"]) & m["区分_9F"].isin(["FCT", "backup"]) & (m["案件確度_9F"] > 0)].copy()

    def st(r):
        if r["_merge"] == "left_only":
            return "照合不可(名称変更/削除)"
        if r["成約"]:
            return "成約"
        if not infy(pd.Series([r["時期"]])).iloc[0]:
            return "期外移動(来期以降)"
        if r["案件確度"] == 0:
            return "0%化(LOST等)"
        if r["区分_9F"] == "backup" and r["区分"] == "FCT":
            return "Backup→FCT昇格"
        if r["区分_9F"] == "FCT" and r["区分"] == "backup":
            return "FCT→Backup降格"
        return f"継続({r['区分']})"

    base["変化"] = base.apply(st, axis=1)
    base["9F区分"] = base["区分_9F"]
    base["モデル"] = base["モデルイニシャル_9F"]
    new = m[(m["_merge"] == "right_only") & infy(m["時期"]) & (m["案件確度"] > 0)].copy()
    new["変化"] = "新規追加"
    new["9F区分"] = "(なし)"
    new["モデル"] = new["モデルイニシャル"]
    allm = pd.concat([base, new])
    tab = pd.crosstab([allm["9F区分"], allm["変化"]], allm["モデル"], margins=True, margins_name="計")
    return allm, tab


# ---------- 代理店 ----------

def agent_tables(d, moved, leads, wl, cfg, as_of):
    p, r = cfg["period"], cfg["rules"]
    op = d[(d["open_FCT"] | d["open_BK"]) & (d["時期"] >= p["h2_start"])].copy()
    pf = d[(d["時期"] > p["fy_end"]) & (d["時期"] <= r["pull_forward_until"]) & d["区分"].isin(["FCT", "backup"])
           & (d["案件確度"] >= r["pull_forward_min_prob"])].copy()
    wr = {m: v["win_rate"] for m, v in cfg["models"].items()}
    op["昇格余地"] = op.apply(lambda x: (wr[x["モデルイニシャル"]] - cfg["backup_rate"]) if x["open_BK"] and x["モデルイニシャル"] in wr else 0, axis=1)
    pf["前倒し余地"] = pf["モデルイニシャル"].map(wr).fillna(0)

    g = op.groupby("一次店")
    t = pd.DataFrame({
        "下期FCT": g["open_FCT"].sum(),
        "下期Backup": g["open_BK"].sum(),
        "着地寄与(台)": g["着地寄与"].sum(),
        "確度加重(台)": g.apply(lambda x: (x["案件確度"] * x["台数"]).sum(), include_groups=False),
        "期日超過": g["期日状態"].apply(lambda s: (s == "期日超過").sum()),
        "アクション未記入": g["期日状態"].apply(lambda s: (s == "アクション未記入").sum()),
        "Backup昇格余地(台)": g["昇格余地"].sum(),
    })
    t["FY27前倒し候補"] = pf.groupby("一次店").size()
    t["前倒し余地(台)"] = pf.groupby("一次店")["前倒し余地"].sum()
    mv = moved[moved["変化"] == "期外移動(来期以降)"]
    t["9F以降の期外移動"] = mv.groupby("一次店").size()
    lg = leads.groupby("一次店")
    t["1H新規リード"] = lg.size()
    t["うちProspecting"] = lg["Prospecting"].sum()
    w = wl[wl["勝敗"].isin(["勝", "敗"])].groupby("一次店")["勝敗"]
    t["FY20-25勝"] = w.apply(lambda s: (s == "勝").sum())
    t["FY20-25敗"] = w.apply(lambda s: (s == "敗").sum())
    t = t.fillna(0)
    n = t["FY20-25勝"] + t["FY20-25敗"]
    t["FY20-25勝率"] = (t["FY20-25勝"] / n).where(n >= 5)
    # 敗け案件は代理店未記入が多く、(未記入)の勝率は意味を持たないため表示しない
    t.loc[t.index.isin(["(未記入)", "未定"]), "FY20-25勝率"] = None
    t["挽回ポテンシャル(台)"] = t["Backup昇格余地(台)"] + t["前倒し余地(台)"]
    tags = op.explode("課題タグ").groupby("一次店")["課題タグ"].agg(
        lambda s: "・".join(f"{k}{v}" for k, v in s.dropna().value_counts().head(3).items()))
    t["主な課題(件数)"] = tags
    t["打ち手区分"] = t.apply(_agent_play, axis=1)
    t = t[(t["下期FCT"] + t["下期Backup"] + t["FY27前倒し候補"] + t["1H新規リード"]) > 0]
    t = t.sort_values(["挽回ポテンシャル(台)", "着地寄与(台)"], ascending=False)
    intcols = ["下期FCT", "下期Backup", "期日超過", "アクション未記入", "FY27前倒し候補", "9F以降の期外移動",
               "1H新規リード", "うちProspecting", "FY20-25勝", "FY20-25敗"]
    t[intcols] = t[intcols].astype(int)
    for c in ["着地寄与(台)", "確度加重(台)", "Backup昇格余地(台)", "前倒し余地(台)", "挽回ポテンシャル(台)"]:
        t[c] = t[c].round(2)
    t["FY20-25勝率"] = t["FY20-25勝率"].round(2)
    t.index.name = "一次店"

    sub = op[op["二次店"] != ""].groupby("二次店").agg(
        案件数=("台数", "size"), 下期FCT=("open_FCT", "sum"), 着地寄与=("着地寄与", "sum"),
        主な一次店=("一次店", lambda s: s.value_counts().index[0]))
    sub = sub.sort_values(["案件数", "着地寄与"], ascending=False).head(20)
    sub["着地寄与"] = sub["着地寄与"].round(2)
    sub["下期FCT"] = sub["下期FCT"].astype(int)
    return t, sub, op, pf


def _agent_play(r):
    tags = []
    if r["下期FCT"] >= 4:
        tags.append("刈り取り集中")
    if r["FY27前倒し候補"] + r["9F以降の期外移動"] >= 3:
        tags.append("前倒し/引き戻し交渉")
    if r["下期Backup"] >= 3:
        tags.append("Backup昇格確認")
    if r["1H新規リード"] >= 5 and r["うちProspecting"] / max(r["1H新規リード"], 1) >= 0.6:
        tags.append("リード案件化テコ入れ")
    if r["期日超過"] + r["アクション未記入"] >= 3:
        tags.append("次アクション設定")
    return "・".join(tags) or "維持"


# ---------- 案件アクションリスト ----------

def action_list(op, pf, cfg, as_of):
    x = pd.concat([op.assign(対象="下期パイプライン"), pf.assign(対象="FY27前倒し候補")])

    def deadline(r):
        g = r["顧客ジャンル"]
        if g == "AC":
            return "2026-11" if r["モデルイニシャル"] in HIGH else "2026-12"
        return "2027-01"

    def play(r):
        tags = r["課題タグ"]
        if r["対象"] == "FY27前倒し候補":
            return "FY26内前倒し可否の確認(年度末予算・補正・価格条件)"
        if r["区分"] == "FCT":
            if r["期日状態"] in ("期日超過", "アクション未記入"):
                return "代理店と次アクション・期日を今週確定"
            if "競合" in tags:
                return "対抗オファー+デモ/リファレンス提示(UTMB型)"
            if "予算・申請待ち" in tags:
                return "予算確定時期の確認・年度末/補正予算の取り込み"
            if "デモ・評価" in tags:
                return "デモ日程確定→評価後2週間で見積・クロージング"
            return "クロージング計画(決裁者・時期・条件)の確認"
        return "FCT昇格条件(予算・時期・仕様)の確認"

    def prio(r):
        if r["対象"] == "下期パイプライン" and r["区分"] == "FCT" and (
                r["時期"] <= 202612 or r["期日状態"] != "期日内"):
            return "A"
        if r["区分"] == "FCT" or r["案件確度"] >= 0.5:
            return "B"
        return "C"

    x["商談期限(目安)"] = x.apply(deadline, axis=1)
    x["推奨打ち手"] = x.apply(play, axis=1)
    x["優先度"] = x.apply(prio, axis=1)
    dfn = cfg["definitions"]
    x["顧客ジャンル"] = x["顧客ジャンル"].map(dfn["genre"]).fillna(x["顧客ジャンル"])
    x["納期確度"] = x["納期確度"].map(lambda v: f"{v}({dfn['delivery'][v]})" if v in dfn["delivery"] else v)
    x["課題分類"] = x["課題タグ"].apply(lambda t: "・".join(t))
    x["案件ID"] = "P1005-R" + x["Excel行"].astype(str)
    x["アクション期日"] = x["期日"].dt.strftime("%Y-%m-%d").fillna("")
    cols = ["優先度", "案件ID", "対象", "担当", "モデルイニシャル", "モデル詳細", "顧客ジャンル", "時期", "区分",
            "案件確度", "納期確度", "一次店", "二次店", "課題分類", "アクション期日", "期日状態", "商談期限(目安)",
            "推奨打ち手"]
    return x[cols].sort_values(["優先度", "時期", "一次店"]).reset_index(drop=True)


# ---------- 本体 ----------

def main(snap_day):
    cfg = tomllib.loads((ROOT / "config.toml").read_text(encoding="utf-8"))
    src = cfg["source"]
    snap = ROOT / "作業" / "snapshots" / snap_day
    as_of = pd.Timestamp(cfg["period"]["as_of"])
    kw = cfg["issue_keywords"]

    a = classify(load_pipeline(snap / src["pipeline_9f"]), cfg)
    b = classify(load_pipeline(snap / src["pipeline_latest"]), cfg)
    ag = b["代理店"].apply(lambda v: split_agent(v, cfg))
    b["一次店"], b["二次店"] = ag.str[0], ag.str[1]
    b["課題タグ"] = b["課題"].apply(lambda t: issue_tags(t, kw))
    b["期日"] = b["アクション時期"].apply(to_date)
    b["期日状態"] = "期日内"
    b.loc[b["期日"] < as_of, "期日状態"] = "期日超過"
    b.loc[b["期日"].isna() & b["アクション"].isna(), "期日状態"] = "アクション未記入"
    b.loc[b["期日"].isna() & b["アクション"].notna(), "期日状態"] = "期日未設定"

    # 前提: 21件の再現と最新化
    land9 = landing(a, cfg, "9FCT時点(再計算)")
    land10 = landing(b, cfg, "10/5時点(最新)")
    deck = DECK_9F.copy()
    deck.loc["計"] = deck.sum()
    deck = deck.reset_index().rename(columns={"index": "モデル"})
    deck.insert(0, "データ", "スライド記載値(9F)")

    moved, mtab = movement(a.copy(), b.copy(), cfg)

    # 新規リード(1H)
    lr = pd.read_excel(snap / src["leads"], sheet_name="1_Salesforce生データ")
    lr = lr[lr["作成日"].notna()].copy()
    lr["一次店"] = lr["一次店"].apply(lambda v: split_agent(v, cfg)[0])
    lr["Prospecting"] = lr["フェーズ"].astype(str).str.startswith("Prospecting")
    lr["フェーズ短縮"] = lr["フェーズ"].astype(str).str.extract(r"^(Prospecting|\d+%)")[0]
    lead_phase = pd.crosstab(lr["一次店"], lr["フェーズ短縮"], margins=True, margins_name="計").sort_values("計", ascending=False)
    lead_fy = pd.crosstab(lr["仕様"], lr["完了予定年度"], margins=True, margins_name="計")

    # 勝敗(FY20-25)
    wl = pd.read_excel(snap / src["winloss"], sheet_name=0, header=None, skiprows=4)
    wl = wl.iloc[:, :12]
    wl.columns = ["担当", "時期", "機種", "モデル", "施設", "代理店", "二次代理店", "ジャンル", "県", "エリア", "デモ", "勝敗"]
    wl = wl[wl["勝敗"].notna()].copy()
    wl["勝敗"] = wl["勝敗"].astype(str).str.strip()
    wl["一次店"] = wl["代理店"].apply(lambda v: split_agent(v, cfg)[0])
    wld = wl[wl["勝敗"].isin(["勝", "敗"])]
    demo = pd.crosstab([wld["機種"]], [wld["デモ"].astype(str).str.strip(), wld["勝敗"]])
    demo_rows = []
    for mdl in ["SH", "MA", "SA", "ID", "FP"]:
        r = {"機種": mdl}
        for dm in ["有", "無"]:
            w = demo.get((dm, "勝"), pd.Series(dtype=int)).get(mdl, 0)
            l = demo.get((dm, "敗"), pd.Series(dtype=int)).get(mdl, 0)
            r[f"デモ{dm}_勝-敗"] = f"{w}-{l}"
            r[f"デモ{dm}_勝率"] = round(w / (w + l), 2) if (w + l) >= 3 else None
        demo_rows.append(r)
    demo_t = pd.DataFrame(demo_rows)
    demo_t = demo_t[demo_t[["デモ有_勝-敗", "デモ無_勝-敗"]].ne("0-0").any(axis=1)]

    agents, sub, op, pf = agent_tables(b, moved, lr, wl, cfg, as_of)
    acts = action_list(op, pf, cfg, as_of)

    # 下期FCTの月別集中
    h2 = b[b["open_FCT"] & (b["時期"] >= cfg["period"]["h2_start"])]
    month_t = pd.crosstab(h2["時期"], h2["納期確度"].fillna("未記入"), margins=True, margins_name="計")

    # 施策表
    ms = pd.read_excel(snap / src["measures"], header=None).dropna(how="all")
    ms = ms.dropna(axis=1, how="all")
    measures = [" / ".join(str(v).strip() for v in r if pd.notna(v)) for r in ms.itertuples(index=False)]

    # 不足の埋め方シナリオ(台換算)
    sc = scenario(b, land10, op, pf, moved, cfg)

    stats = dict(
        land9=land9, land10=land10, deck=deck, mtab=mtab, moved=moved, agents=agents, sub=sub, acts=acts,
        lead_phase=lead_phase, lead_fy=lead_fy, lr=lr, demo_t=demo_t, wl=wl, month_t=month_t, measures=measures,
        sc=sc, op=op, pf=pf, b=b, cfg=cfg)
    out = ROOT / "出力"
    out.mkdir(exist_ok=True)
    tag = snap_day.replace("-", "")
    write_excel(out / f"パイプライン不足分析_{tag}.xlsx", stats, cfg)
    write_md(out / f"パイプライン不足分析_{tag}.md", stats, cfg)
    print("written:", out)


def scenario(b, land10, op, pf, moved, cfg):
    tot = land10[land10["モデル"] == "計"].iloc[0]
    short = -float(tot["不足(台)"])
    wr = {m: v["win_rate"] for m, v in cfg["models"].items()}
    bk = op[op["open_BK"] & op["モデルイニシャル"].isin(list(wr))]
    up_unit = sum(wr[m] - cfg["backup_rate"] for m in bk["モデルイニシャル"])
    mv = moved[(moved["変化"] == "期外移動(来期以降)") & moved["モデル"].isin(list(wr))]
    back_unit = sum(wr[m] - 0 for m in mv["モデル"])
    avg_wr = sum(wr.values()) / len(wr)
    sc = cfg["scenario"]
    up, pb, lm, nf = sc["upgrade_rate"], sc["pullback_rate"], sc["lead_meeting_rate"], sc["new_to_fct_rate"]
    tm = sc["telemarketing_targets"] * lm
    gl = sc["guarantee_leads"] * lm
    ex = sc["expo_leads"] * 2 * lm  # NGS EXPO(10月)+細胞凝集研究会(11月)の2回
    rows = [
        ("不足(10/5時点)", None, None, round(short, 2), "着地理論値 − 9F目標(モデル合計)"),
        ("① 下期Backup→FCT昇格", len(bk), f"昇格{up:.0%}", round(up_unit * up, 2),
         f"対象{len(bk)}件 × 昇格率{up:.0%} × (モデル勝率−10%)"),
        ("② 期外移動案件の引き戻し", len(mv), f"引き戻し{pb:.0%}", round(back_unit * pb, 2),
         f"9F以降に来期へ移った{len(mv)}件 × 引き戻し率{pb:.0%} × モデル勝率"),
        (f"③ テレマ({sc['telemarketing_targets']}名)からの新規FCT", round(tm, 1), f"商談化{lm:.0%}・FCT化{nf:.0%}",
         round(tm * nf * avg_wr, 2),
         f"{sc['telemarketing_targets']}名 × 商談化{lm:.0%}≒{tm:.1f}商談(slide7保守ケース相当) × FCT化{nf:.0%}(slide4: 30/94) × 平均勝率{avg_wr:.2f}"),
        (f"④ 150リード保証({sc['guarantee_leads']}件)からの新規FCT", round(gl, 1), f"商談化{lm:.0%}・FCT化{nf:.0%}",
         round(gl * nf * avg_wr, 2), f"{sc['guarantee_leads']}件 × {lm:.0%}≒{gl:.1f}商談 × {nf:.0%} × {avg_wr:.2f}"),
        ("⑤ 展示会・学会(NGS EXPO/細胞凝集研究会)", round(ex, 1), f"商談化{lm:.0%}・FCT化{nf:.0%}",
         round(ex * nf * avg_wr, 2),
         f"{sc['expo_leads']}リード(FY25-Q2展示会実績)×2回 × {lm:.0%}≒{ex:.1f}商談 × {nf:.0%} × {avg_wr:.2f}"),
    ]
    t = pd.DataFrame(rows, columns=["項目", "対象件数", "率", "台", "算出根拠"])
    filled = t.loc[1:, "台"].sum()
    t.loc[len(t)] = ["①〜⑤ 合計", None, None, round(filled, 2), "①〜⑤の和"]
    t.loc[len(t)] = ["残ギャップ", None, None, round(short - filled, 2), "不足 − 合計(正なら未充足)"]
    agl = up_unit * up + back_unit * pb
    if short - filled > 0 and agl > 0:
        k = (short - (filled - agl)) / agl
        t.loc[len(t)] = ["参考: 不足を埋める昇格率・引き戻し率", None, f"昇格{up * k:.0%}・引き戻し{pb * k:.0%}", None,
                         f"インハウス施策は据え置き、①②の率を同じ倍率({k:.1f}倍)で引き上げた場合"]
    return t


# ---------- 書き出し ----------

def write_excel(path, s, cfg):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    sheets = [
        ("00_前提と算出式", premise_df(s, cfg)),
        ("01_21件再現と最新化", pd.concat([s["deck"], s["land9"], s["land10"]], ignore_index=True)),
        ("02_9F→10月の案件移動", s["mtab"].reset_index()),
        ("03_代理店別(一次店)", s["agents"].reset_index()),
        ("04_販売店別(二次店)上位", s["sub"].reset_index()),
        ("05_挽回シナリオ", s["sc"]),
        ("06_新規リード×代理店", s["lead_phase"].reset_index()),
        ("07_新規リード完了予定年度", s["lead_fy"].reset_index()),
        ("08_デモ有無と勝率", s["demo_t"]),
        ("09_下期FCT月別", s["month_t"].reset_index()),
        ("10_案件アクションリスト", s["acts"]),
        ("11_短期アクション計画", plan_df(s)),
    ]
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets:
            df.to_excel(xw, sheet_name=name[:31], index=False)
            ws = xw.sheets[name[:31]]
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="305496")
                c.alignment = Alignment(wrap_text=True, vertical="center")
            for j, col in enumerate(df.columns, 1):
                vals = [str(col)] + [str(v) for v in df.iloc[:200, j - 1].tolist()]
                width = min(60, max(8, max(len(v) for v in vals) * 1.6))
                ws.column_dimensions[get_column_letter(j)].width = width
            ws.freeze_panes = "B2"


def premise_df(s, cfg):
    m = cfg["models"]
    rows = [
        ("データソース", "案件Excel(最新)", f"{cfg['source']['pipeline_latest']} / シート MOC Pipeline 1005", "確定"),
        ("データソース", "案件Excel(9F)", f"{cfg['source']['pipeline_9f']} / シート MOC Pipeline 9FCT", "確定"),
        ("データソース", "施策表", f"{cfg['source']['measures']} / Sheet1", "確定"),
        ("データソース", "新規リード", f"{cfg['source']['leads']} / 1_Salesforce生データ", "確定"),
        ("データソース", "勝敗実績", f"{cfg['source']['winloss']} / FY20_FY25商談統括", "確定"),
        ("取り込み", "スナップショット", "原本/ → 作業/snapshots/<日付>/ にコピーし、コピーのみ読込。SHA-256で原本不変を確認", "確定"),
        ("期間", "今期", "FY26 = 2026/04〜2027/03、下期 = 2026/10〜", "確定(資料表記より)"),
        ("算出式", "着地理論値", "成約 + FCT案件(未成約)×モデル勝率 + Backup×10%", "確定(Sep資料 slide2)"),
        ("算出式", "必要リード", "不足台数 ÷ モデル勝率 をモデル別に計算し合計(SH6+MA6+ID7+FP2=21件)", "確定(Sep資料 slide2)"),
        ("算出式", "端数", "スライドはモデル別に丸め(MA 6.7→6, ID 6.7→7)。本分析は小数のまま表示", "注記"),
        ("確度区分", "FCT", "今期FCST計上案件。モデル勝率で加重: " + ", ".join(f"{k}{v['win_rate']:.0%}" for k, v in m.items()), "確定(Sep資料)"),
        ("確度区分", "Backup", "FCT外の今期案件。一律10%で加重", "確定(Sep資料)"),
        ("確度区分", "成約", "区分列が「nFCT」(例: 9FCT=9月FCTで成約)", "データから推定"),
        ("確度区分", "0%", "Sales Prob 0% は LOST/休止として集計外", "確定(slide4 注記)"),
        ("確度区分", "案件確度(Sales Prob %)", "代理店別の『確度加重(台)』でのみ参考利用", "参考"),
        ("列対応", "案件Excel", "見出しと実データが1列ずれているため 'Sales month' 列基準の位置で割当(担当=左隣)", "データから推定"),
        ("列対応", "顧客ジャンル", "AC=アカデミア, BT=バイオテック企業, PHM=製薬企業", "確定(10/6確認)"),
        ("列対応", "顧客ジャンル BTL", "定義外のため BT(バイオテック企業)扱い", "仮置き"),
        ("列対応", "納期確度", "A=月決, B=Q決, C=未定", "確定(10/6確認)"),
        ("挽回シナリオ", "率", "Backup昇格15%・引き戻し10%・リード商談化2%(厳しめ設定)", "確定(10/6確認)"),
        ("代理店", "一次店の判定", "主要一次店(" + "・".join(cfg["rules"]["primary_agents"]) + ")を含めばそれを一次店、無ければ末尾を一次店", "仮置き"),
        ("代理店", "BTL", "バイオテック・ラボの略と解釈", "仮置き"),
        ("上期▲7台", "扱い", "9F目標60台に含まれる(1H実績11台+下期49台)。9F達成=上期▲7台の挽回", "確定(10/6確認)"),
        ("商談期限", "目安", "High(ID/FP): アカデミア11月・企業1月 / Low-Mid(SH/MA/SA): アカデミア12月・企業1月", "確定(Sep資料 slide3)"),
        ("個人情報", "出力方針", "施設名・個人名・自由記述原文は出力しない。案件はスナップショットのExcel行番号(案件ID)で参照", "確定"),
    ]
    return pd.DataFrame(rows, columns=["区分", "項目", "内容", "確度"])


def plan_df(s):
    return pd.DataFrame(PLAN_ROWS(s), columns=["週", "期間", "軸", "アクション", "対象・数", "KPI(週末時点)", "根拠"])


def PLAN_ROWS(s):
    ag = s["agents"]
    acts = s["acts"]
    nA = int((acts["優先度"] == "A").sum())
    top = [a for a in ag.index if a not in ("(未記入)", "未定")][:5]
    plan_m = list(PLAN_MODELS)
    s_up = s["cfg"]["scenario"]["upgrade_rate"]
    op = s["op"]
    nbk = int((op["open_BK"] & op["モデルイニシャル"].isin(plan_m)).sum())
    mvr = s["moved"][(s["moved"]["変化"] == "期外移動(来期以降)") & s["moved"]["モデル"].isin(plan_m)]
    pfr = s["pf"][s["pf"]["モデルイニシャル"].isin(plan_m)]
    nmv = len(mvr)
    nunion = len(set(mvr["Excel行"].dropna().astype(int)) | set(pfr["Excel行"].astype(int)))
    stale = int(ag["期日超過"].sum() + ag["アクション未記入"].sum())
    return [
        ("W1", "10/6-10/10", "代理店", "優先度A案件の次アクション・期日を代理店と確定(電話/同行)",
         f"A案件{nA}件", "A案件の期日設定率100%", "10_案件アクションリスト"),
        ("W1", "10/6-10/10", "代理店", f"上位一次店({'・'.join(top)})と案件棚卸会(30分/社)",
         f"{len(top)}社", "Backup昇格/前倒し候補の可否を全件判定", "03_代理店別 挽回ポテンシャル上位"),
        ("W1", "10/6-10/10", "インハウス", "テレマ開始(10/9〜) S/A 44名を先行架電、営業へ即日トス",
         "S/A 44名→全422名", "接触率・商談化件数を日次集計", "Sep資料 slide6-7, 施策表"),
        ("W1", "10/6-10/10", "全体", "期日超過・アクション未記入の案件を担当別に解消",
         f"{stale}件", "未記入0件", "03_代理店別"),
        ("W2", "10/13-10/17", "代理店", "Backup→FCT昇格確認(予算・時期・仕様の3点)",
         f"下期Backup{nbk}件", f"昇格判定完了、昇格{s_up:.0%}以上", "05_挽回シナリオ①"),
        ("W2", "10/13-10/17", "代理店", "来期へ移った案件の引き戻し交渉(年度末予算・補正・価格条件)",
         f"{nunion}件(9F以降の期外移動{nmv}件とFY27前倒し候補の和集合)", "引き戻し可否の回答回収", "02_案件移動, 05_挽回シナリオ②"),
        ("W2", "10/13-10/17", "インハウス", "NGS EXPO2026(10月)リードを48h以内に一次店へ配分・同行設定",
         "展示会リード全件", "48h以内配分率", "施策表"),
        ("W3", "10/20-10/24", "代理店", "High(ID/FP)アカデミア案件: 11月期限に向けデモ/見積を完了",
         "ID/FP FCT・Backup", "見積提出・デモ日程確定", "Sep資料 slide3 商談期限, 08_デモ有無"),
        ("W3", "10/20-10/24", "インハウス", "1H新規リードのProspecting滞留を代理店別に案件化判定",
         f"Prospecting {int(s['lr']['Prospecting'].sum())}件", "フェーズ更新率50%", "06_新規リード×代理店"),
        ("W3", "10/20-10/24", "代理店", "主力販売店(二次店)向け 年度末予算案件の掘り起こし依頼(事例資料付)",
         "04_販売店上位", "新規案件の登録数", "Sep資料 slide3 User Voice"),
        ("W4", "10/27-10/31", "全体", "10月末FCTレビュー: 着地理論値を再計算し不足と必要件数を更新",
         "全モデル", "不足台数の縮小幅", "01_21件再現と最新化"),
        ("W4", "10/27-10/31", "インハウス", "テレマ結果(商談化率・単価)評価→11月の架電対象/予算を決定",
         "422名", "商談化率 2-5%", "Sep資料 slide7"),
        ("W4", "10/27-10/31", "インハウス", "細胞凝集研究会(11月)の事前アポ設定(既存パイプライン施設を優先)",
         "来場見込み施設", "事前アポ件数", "施策表"),
    ]


def write_md(path, s, cfg):
    l9, l10 = s["land9"], s["land10"]
    sc = s["sc"]
    t9 = l9[l9["モデル"] == "計"].iloc[0]
    t10 = l10[l10["モデル"] == "計"].iloc[0]
    mv = s["moved"]
    mvp = mv[mv["モデル"].isin(list(PLAN_MODELS))]
    nmove = int((mvp["変化"] == "期外移動(来期以降)").sum())
    nmove_bk = int(((mvp["変化"] == "期外移動(来期以降)") & (mvp["9F区分"] == "backup")).sum())
    nnew = int((mvp["変化"] == "新規追加").sum())
    inh = float(sc.iloc[3:6]["台"].sum())
    agl = float(sc.iloc[1:3]["台"].sum())
    short = float(sc.iloc[0]["台"])
    demo = s["demo_t"].set_index("機種")
    ag = s["agents"]
    ag_md = ag[["下期FCT", "下期Backup", "着地寄与(台)", "FY27前倒し候補", "9F以降の期外移動", "1H新規リード",
                "うちProspecting", "FY20-25勝率", "期日超過", "アクション未記入", "挽回ポテンシャル(台)", "打ち手区分"]].head(15)
    lr = s["lr"]
    pros = int(lr["Prospecting"].sum())
    gap_row = sc[sc["項目"] == "残ギャップ"].iloc[0]
    gap_left = float(gap_row["台"])
    need_row = sc[sc["項目"].str.startswith("参考")]
    need_rates = need_row.iloc[0]["率"] if len(need_row) else ""
    acts = s["acts"]
    nA = int((acts["優先度"] == "A").sum())
    known = ag.drop(index=[i for i in ("(未記入)",) if i in ag.index])
    known = known.drop(index=[i for i in ("未定",) if i in known.index])
    top5 = known["着地寄与(台)"].nlargest(5)
    share_top = top5.sum() / max(ag["着地寄与(台)"].sum(), 1e-9)
    noagent = ag.loc["(未記入)"] if "(未記入)" in ag.index else None
    m10 = s["month_t"]
    mar = int(m10.loc[202703, "計"]) if 202703 in m10.index else 0
    h2tot = int(m10.loc["計", "計"])

    L = []
    L.append(f"# パイプライン不足分析と短期アクション(基準日 {cfg['period']['as_of']})\n")
    L.append("> 出典: 作業/snapshots/2026-10-06/ のコピー(原本は未変更、MANIFEST.sha256で照合)。"
             "施設名・個人名・自由記述の原文は載せていません。案件は `案件ID`(P1005-R<スナップショットのExcel行>)で参照してください。"
             "**仮置き**と書いた数字は前提が確定していないものです。\n")

    L.append("## 1. 結論\n")
    L.append(f"1. **21件は9FCT時点の数字で、10/5時点では約{t10['必要リード(件)']:.0f}件(FCT相当)に増えている。** "
             f"着地理論値は {t9['着地理論値']:.1f}台(9F)から {t10['着地理論値']:.1f}台(10/5)に下がり、"
             f"9F目標60台(1H実績11台+下期49台。達成すれば上期▲7台も挽回)に対する不足は {-t9['不足(台)']:.1f}台 → {-t10['不足(台)']:.1f}台に広がった。"
             f"主な原因は、9F以降に今期の案件{nmove}件(うちBackup {nmove_bk}件)が来期以降へ移ったこと。今期の新規追加は{nnew}件にとどまる。")
    L.append(f"2. **埋め手の主力は代理店がすでに持っている案件。インハウスリードは補助。** "
             f"不足{short:.1f}台に対し、Backup昇格と期外移動案件の引き戻しで{agl:.1f}台、"
             f"テレマ・リード保証・展示会で{inh:.1f}台(厳しめの率で試算、表6)。"
             + (f"合計{agl + inh:.1f}台で不足をかろうじて上回る程度なので、率が下振れすれば届かない。"
                if gap_left <= 0 else f"合計{agl + inh:.1f}台で、まだ{gap_left:.1f}台足りない。埋めるには、代理店との棚卸で{need_rates}を実現するか、新規の大型案件を取る必要がある。")
             + 
             "新規登録から今期FCTになるのは約32%(slide4: 30/94)。商談期限(アカデミアはHighが11月・Low-Midが12月、企業は1月)にも間に合わせる必要がある。")
    L.append(f"3. **下期の着地寄与は上位5社({'・'.join(top5.index)})で{share_top:.0%}を占める。** "
             "表5の『挽回ポテンシャル』が大きい順に代理店と案件棚卸会を開き、Backupを昇格できるか・前倒しできるかを今週〜来週で全件判断する。")
    if noagent is not None:
        L.append(f"4. **代理店未記入の案件が下期に{int(noagent['下期FCT'] + noagent['下期Backup'])}件、9F以降の期外移動にも{int(noagent['9F以降の期外移動'])}件ある。** "
                 "『9割が代理店経由』という前提を、データ上確かめられない部分。担当者が一次店・二次店を記入し、代理店別の精度を上げる。")
    L.append(f"5. **1H新規リード{len(lr)}件のうち{pros}件({pros / len(lr):.0%})がProspectingのまま。** "
             "量は1H目標(140件)を上回っているので、ボトルネックは獲得量ではなく案件化(フェーズ更新・代理店への引き渡し)。")
    if "ID" in demo.index and demo.loc["ID", "デモ有_勝率"] is not None:
        L.append(f"6. **IDはデモの有無で勝率が大きく違う(デモ有 {demo.loc['ID', 'デモ有_勝-敗']}={demo.loc['ID', 'デモ有_勝率']:.0%}、"
                 f"デモ無 {demo.loc['ID', 'デモ無_勝-敗']}={demo.loc['ID', 'デモ無_勝率']:.0%}、FY20-25)。** "
                 "IDは必要リードが最も多い(勝率30%)ので、下期のID案件はデモを必須にし、11月の期限(アカデミア)から逆算して日程を押さえる。SH/MA/SAではデモ有無の差はほとんどない。")
    mar_c = int(m10.loc[202703, "C"]) if (202703 in m10.index and "C" in m10.columns) else 0
    h2_c = int(m10.loc["計", "C"]) if "C" in m10.columns else 0
    L.append(f"7. **下期FCT {h2tot}件のうち{mar}件が2027年3月に集中し、そのうち{mar_c}件は納期確度C(未定)。** "
             f"下期FCT全体でも{h2_c}件が納期未定のまま。来期へずれ込むリスクが高いので、棚卸で納期をB(Q決)以上に固められる案件と、11〜12月に前倒しできる案件を見極める。\n")

    L.append("## 2. 前提と算出式\n")
    L.append("- 着地理論値 = 成約 + FCT案件(未成約)×モデル勝率 + Backup×10%(Sep資料 slide2)")
    L.append("- 必要リード = 不足台数 ÷ モデル勝率(モデル別に計算して合計)。スライドの21件 = SH6 + MA6 + ID7 + FP2(端数はモデルごとに丸め)")
    L.append("- 勝率: " + "、".join(f"{k} {v['win_rate']:.0%}" for k, v in cfg["models"].items()) + "(スライド記載値。FY20-25の実勝率74%とは定義が違う)")
    L.append("- 上期▲7台: 9F目標60台(1H実績11台 + 下期49台)に含まれる。9F目標を達成すれば上期▲7台も挽回できる(10/6確認)")
    L.append("- 顧客ジャンル: AC=アカデミア / BT=バイオテック企業 / PHM=製薬企業(10/6確認)。BTL(2件)は定義外のためBT扱い(**仮置き**)")
    L.append("- 納期確度: A=月決 / B=Q決 / C=未定(10/6確認)")
    sc_cfg = cfg["scenario"]
    L.append(f"- 挽回シナリオの率: Backup昇格{sc_cfg['upgrade_rate']:.0%}・引き戻し{sc_cfg['pullback_rate']:.0%}・リード商談化{sc_cfg['lead_meeting_rate']:.0%}(厳しめ設定、10/6確認)")
    L.append("- 代理店表記の BTL はバイオテック・ラボの略と解釈(**仮置き**)")
    L.append("- 案件Excelは見出しと実データが1列ずれているので、'Sales month'列を基準に列の位置で読んでいる\n")

    L.append("## 3. 21件の再現と最新化(モデル別)\n")
    L.append("スライドの値・9Fデータでの再計算・10/5データでの再計算を並べた。9Fでの再計算はスライドとほぼ一致する(MA・IDでFCT件数が±1違う)。\n")
    show = ["データ", "モデル", "実績(成約)", "FCT案件", "Backup", "着地理論値", "不足(台)", "必要リード(件)"]
    L.append(md_table(pd.concat([s["deck"].rename(columns={"FCST案件": "FCT案件", "9F不足": "不足(台)", "必要リード": "必要リード(件)"}),
                                 l9, l10], ignore_index=True).reindex(columns=show)))
    L.append("")

    L.append("## 4. 9F→10/5の案件移動\n")
    L.append("9F時点で今期にあったFCT/Backup案件(確度>0)が、10/5時点でどうなったか(施設名+モデル詳細で照合)。\n")
    L.append(md_table(s["mtab"], index=True))
    L.append("")

    L.append("## 5. 代理店別(一次店)\n")
    L.append("- 着地寄与 = FCT×モデル勝率 + Backup×10%(下期・確度>0)")
    L.append("- 挽回ポテンシャル = 下期Backupの昇格余地(モデル勝率−10%) + FY27前倒し候補(2027/4〜9月・確度40%以上)×モデル勝率。案件を全部取れた場合の上限値")
    L.append("- FY20-25勝率は勝+敗が5件以上ある代理店だけ表示\n")
    L.append(md_table(ag_md, index=True))
    L.append("")
    L.append("**読み取り**\n")
    for name, r in ag.head(6).iterrows():
        if name in ("(未記入)", "未定"):
            continue
        L.append(f"- {name}: 下期FCT{r['下期FCT']}件・Backup{r['下期Backup']}件、FY27前倒し候補{r['FY27前倒し候補']}件、"
                 f"9F以降の期外移動{r['9F以降の期外移動']}件、未対応(期日超過+アクション未記入){r['期日超過'] + r['アクション未記入']}件"
                 + (f"、FY20-25勝率{r['FY20-25勝率']:.0%}" if pd.notna(r['FY20-25勝率']) else "") + f" → {r['打ち手区分']}")
    L.append("- FY20-25の勝率は、敗け案件で代理店が未記入のことが多いため、実際より高めに出ている可能性がある")
    L.append("")
    L.append("### 販売店(二次店)上位\n")
    L.append(md_table(s["sub"], index=True))
    L.append("")

    L.append("## 6. 不足の埋め方(シナリオ)\n")
    L.append(md_table(sc))
    L.append("\n率は厳しめの設定(10/6確認)。W2の棚卸で昇格・引き戻しの実績が出たら、その値に置き換えて再計算する。\n")

    L.append("## 7. インハウスリードと施策\n")
    L.append("直近の施策(施策表): " + " ／ ".join(s["measures"][1:]) + "\n")
    L.append("FY20-25のデモ有無別勝率(勝敗が3件以上の機種のみ):\n")
    L.append(md_table(s["demo_t"]))
    L.append("\n新規リードのフェーズ(一次店別、上位):\n")
    L.append(md_table(s["lead_phase"].head(12), index=True))
    L.append("")

    L.append("## 8. 短期アクション計画(10/6〜10/31、4週間)\n")
    L.append(md_table(plan_df(s)))
    L.append(f"\n優先度Aの案件は{nA}件(下期FCTで、時期が12月以前か、期日超過・未記入・期日未設定のもの)。"
             "全件のリストはExcelの「10_案件アクションリスト」シートにある。\n")

    L.append("## 9. 未確定事項(確認したいこと)\n")
    L.append("- 顧客ジャンル『BTL』(2件)の意味(本分析ではBT=バイオテック企業として扱った)、『CL』(前倒し候補に1件)の意味")
    L.append("- 代理店表記『BTL』はバイオテック・ラボの略でよいか")
    L.append("- 『9割が代理店紹介』の定義(紹介元か、納入経路か)。データ上は代理店未記入の案件がある")
    Path(path).write_text("\n".join(L) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "2026-10-06")
